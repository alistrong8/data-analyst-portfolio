import pandas as pd
import xlsxwriter
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
import os

BASE_DIR = "/Users/mac/Desktop/App/portfolio/data-analyst-portfolio"
DATA_DIR = f"{BASE_DIR}/Datasets"
EXCEL_DIR = f"{BASE_DIR}/Excel_Projects"

def create_sales_dashboard():
    print("Creating Sales Performance Dashboard...")
    df = pd.read_csv(f"{DATA_DIR}/ecommerce_sales.csv")
    
    # Create Excel Writer
    writer = pd.ExcelWriter(f"{EXCEL_DIR}/Sales_Performance_Dashboard.xlsx", engine='xlsxwriter')
    workbook = writer.book
    
    # 1. Raw Data Sheet
    df.to_excel(writer, sheet_name='Raw_Data', index=False)
    worksheet_data = writer.sheets['Raw_Data']
    worksheet_data.set_column('A:J', 15)
    
    # 2. Summary Sheet (Pivot-like)
    # Monthly Revenue
    df['Date'] = pd.to_datetime(df['Date'])
    monthly_sales = df.groupby(df['Date'].dt.to_period('M'))['TotalPrice'].sum().reset_index()
    monthly_sales.columns = ['Month', 'Revenue']
    monthly_sales['Month'] = monthly_sales['Month'].dt.strftime('%Y-%m')
    
    monthly_sales.to_excel(writer, sheet_name='Dashboard_Data', startcol=0, startrow=0, index=False)
    
    # Category Sales
    cat_sales = df.groupby('Category')['TotalPrice'].sum().reset_index()
    cat_sales.to_excel(writer, sheet_name='Dashboard_Data', startcol=4, startrow=0, index=False)
    
    dashboard_sheet = writer.sheets['Dashboard_Data']
    
    # Create Formats
    currency_fmt = workbook.add_format({'num_format': '$#,##0'})
    bold_fmt = workbook.add_format({'bold': True})
    
    dashboard_sheet.set_column('B:B', 15, currency_fmt)
    dashboard_sheet.set_column('E:E', 15, currency_fmt)
    
    # 3. Create Charts
    # Line Chart: Monthly Trend
    chart_trend = workbook.add_chart({'type': 'line'})
    chart_trend.add_series({
        'name':       'Monthly Revenue',
        'categories': ['Dashboard_Data', 1, 0, len(monthly_sales), 0],
        'values':     ['Dashboard_Data', 1, 1, len(monthly_sales), 1],
    })
    chart_trend.set_title({'name': 'Monthly Revenue Trend'})
    chart_trend.set_x_axis({'name': 'Month'})
    chart_trend.set_y_axis({'name': 'Revenue ($)'})
    dashboard_sheet.insert_chart('A15', chart_trend)
    
    # Column Chart: Category Performance
    chart_cat = workbook.add_chart({'type': 'column'})
    chart_cat.add_series({
        'name':       'Revenue by Category',
        'categories': ['Dashboard_Data', 1, 4, len(cat_sales), 4],
        'values':     ['Dashboard_Data', 1, 5, len(cat_sales), 5],
        'data_labels': {'value': True}
    })
    chart_cat.set_title({'name': 'Sales by Category'})
    dashboard_sheet.insert_chart('E15', chart_cat)
    
    writer.close()
    print(f"Saved {EXCEL_DIR}/Sales_Performance_Dashboard.xlsx")

def create_financial_model():
    print("Creating Financial Forecasting Model...")
    df = pd.read_csv(f"{DATA_DIR}/financial_forecast.csv")
    
    writer = pd.ExcelWriter(f"{EXCEL_DIR}/Financial_Forecasting_Model.xlsx", engine='xlsxwriter')
    workbook = writer.book
    
    # 1. Data & Model
    # We will write the data, but add formula columns for "Variance"
    
    sheet_name = 'Financial_Model'
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    worksheet = writer.sheets[sheet_name]
    
    # Add Variance Column (Budget vs Actual) using Formulas
    # Assuming columns: A=Month, B=Revenue, F=BudgetedRevenue
    # Variance = Revenue - BudgetedRevenue
    
    # Write Header
    header_fmt = workbook.add_format({'bold': True, 'bottom': 1})
    worksheet.write(0, 6, "Variance (Act vs Bud)", header_fmt)
    worksheet.write(0, 7, "Variance %", header_fmt)
    
    # Write Formulas
    money_fmt = workbook.add_format({'num_format': '$#,##0.00'})
    pct_fmt = workbook.add_format({'num_format': '0.0%'})
    
    for row_num in range(1, len(df) + 1):
        # Revenue is col B (index 1), Budget is col F (index 5)
        # Excel rows are 1-based, but here we iterate 1 to N.
        # So row 2 in Excel is index 1 here.
        excel_row = row_num + 1
        
        # Variance = B{row} - F{row}
        worksheet.write_formula(row_num, 6, f'=B{excel_row}-F{excel_row}', money_fmt)
        
        # Variance % = (B{row} - F{row}) / F{row}
        worksheet.write_formula(row_num, 7, f'=(B{excel_row}-F{excel_row})/F{excel_row}', pct_fmt)
        
    worksheet.set_column('A:H', 15)
    
    # Conditional Formatting for Variance
    worksheet.conditional_format(1, 6, len(df), 6, {
        'type': 'cell',
        'criteria': '<',
        'value': 0,
        'format': workbook.add_format({'font_color': '#9C0006', 'bg_color': '#FFC7CE'})
    })
    
    worksheet.conditional_format(1, 6, len(df), 6, {
        'type': 'cell',
        'criteria': '>',
        'value': 0,
        'format': workbook.add_format({'font_color': '#006100', 'bg_color': '#C6EFCE'})
    })

    writer.close()
    print(f"Saved {EXCEL_DIR}/Financial_Forecasting_Model.xlsx")

def create_hr_dashboard():
    print("Creating HR Attrition Dashboard...")
    df = pd.read_csv(f"{DATA_DIR}/hr_employee_attrition.csv")
    
    writer = pd.ExcelWriter(f"{EXCEL_DIR}/HR_Attrition_Dashboard.xlsx", engine='xlsxwriter')
    workbook = writer.book
    
    df.to_excel(writer, sheet_name='HR_Data', index=False)
    
    # Analysis: Attrition by Department
    attrition_dept = df[df['Attrition']=='Yes'].groupby('Department').size().reset_index(name='AttritionCount')
    attrition_dept.to_excel(writer, sheet_name='Analysis', index=False)
    
    # Analysis: Avg Satisfaction by Attrition
    sat_analysis = df.groupby('Attrition')['JobSatisfaction'].mean().reset_index()
    sat_analysis.to_excel(writer, sheet_name='Analysis', startcol=4, index=False)
    
    analysis_sheet = writer.sheets['Analysis']
    
    # Create Chart: Attrition by Dept
    chart = workbook.add_chart({'type': 'column'})
    chart.add_series({
        'name': 'Attrition Count',
        'categories': ['Analysis', 1, 0, len(attrition_dept), 0],
        'values': ['Analysis', 1, 1, len(attrition_dept), 1],
    })
    chart.set_title({'name': 'Attrition by Department'})
    analysis_sheet.insert_chart('A10', chart)
    
    writer.close()
    print(f"Saved {EXCEL_DIR}/HR_Attrition_Dashboard.xlsx")

def create_inventory_dashboard():
    print("Creating Inventory Dashboard...")
    df = pd.read_csv(f"{DATA_DIR}/supply_chain_inventory.csv")
    
    writer = pd.ExcelWriter(f"{EXCEL_DIR}/Inventory_Management.xlsx", engine='xlsxwriter')
    workbook = writer.book
    
    df.to_excel(writer, sheet_name='Inventory_Data', index=False)
    worksheet = writer.sheets['Inventory_Data']
    
    # Conditional Formatting for Stockout Risk
    # Column G is StockoutRisk (7th col, index 6)
    worksheet.conditional_format(1, 6, len(df), 6, {
        'type': 'text',
        'criteria': 'containsText',
        'value': 'High',
        'format': workbook.add_format({'bg_color': '#FFC7CE', 'font_color': '#9C0006'})
    })
    
    writer.close()
    print(f"Saved {EXCEL_DIR}/Inventory_Management.xlsx")

if __name__ == "__main__":
    create_sales_dashboard()
    create_financial_model()
    create_hr_dashboard()
    create_inventory_dashboard()
